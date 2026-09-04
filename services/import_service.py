# services/import_service.py

from pathlib import Path

from openpyxl import load_workbook

from database.repository import Repository


class ImportService:
    REQUIRED_COLUMNS = {
        "модель": "model",
        "серийный номер": "serial_number",
        "гаражный номер": "garage_number",
    }

    OPTIONAL_COLUMNS = {
        "госномер": "registration_number",
        "гос. номер": "registration_number",
        "государственный номер": "registration_number",
        "год выпуска": "manufacture_year",
        "наработка": "current_hours",
        "статус": "status",
        "примечание": "note",
    }

    def __init__(self, repository=None):
        self.repository = repository or Repository()

    def import_equipment(self, file_path):
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"Файл не найден: {path}"
            )

        if path.suffix.lower() != ".xlsx":
            raise ValueError(
                "Поддерживается импорт только из файлов .xlsx"
            )

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        try:
            worksheet = workbook.active

            header_row = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                None,
            )

            if not header_row:
                raise ValueError(
                    "Файл не содержит заголовков."
                )

            column_map = self._build_column_map(
                header_row
            )

            self._validate_required_columns(
                column_map
            )

            result = {
                "imported": 0,
                "skipped": 0,
                "errors": [],
            }

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=2,
                    values_only=True,
                ),
                start=2,
            ):
                try:
                    data = self._read_row(
                        row,
                        column_map,
                    )

                    if self._is_empty_row(data):
                        continue

                    self._validate_equipment_row(
                        data,
                        row_number,
                    )

                    if self.repository.equipment_exists(
                        model=data["model"],
                        serial_number=data["serial_number"],
                        garage_number=data["garage_number"],
                    ):
                        result["skipped"] += 1
                        continue

                    self.repository.add_equipment(
                        model=data["model"],
                        serial_number=data["serial_number"],
                        garage_number=data["garage_number"],
                        registration_number=data[
                            "registration_number"
                        ],
                        manufacture_year=data[
                            "manufacture_year"
                        ],
                        current_hours=data[
                            "current_hours"
                        ],
                        status=data["status"],
                        note=data["note"],
                    )

                    result["imported"] += 1

                except Exception as error:
                    result["errors"].append(
                        f"Строка {row_number}: {error}"
                    )

            return result

        finally:
            workbook.close()

    def _build_column_map(self, header_row):
        column_map = {}

        all_columns = {
            **self.REQUIRED_COLUMNS,
            **self.OPTIONAL_COLUMNS,
        }

        for index, value in enumerate(
            header_row
        ):
            normalized = self._normalize_header(
                value
            )

            field_name = all_columns.get(
                normalized
            )

            if field_name:
                column_map[field_name] = index

        return column_map

    def _validate_required_columns(
        self,
        column_map,
    ):
        missing = []

        for header, field_name in (
            self.REQUIRED_COLUMNS.items()
        ):
            if field_name not in column_map:
                missing.append(header)

        if missing:
            raise ValueError(
                "Отсутствуют обязательные столбцы: "
                + ", ".join(missing)
            )

    def _read_row(
        self,
        row,
        column_map,
    ):
        def value(field_name, default=""):
            index = column_map.get(field_name)

            if index is None:
                return default

            if index >= len(row):
                return default

            cell_value = row[index]

            if cell_value is None:
                return default

            return cell_value

        manufacture_year = self._to_int_or_none(
            value(
                "manufacture_year",
                None,
            )
        )

        current_hours = self._to_float(
            value(
                "current_hours",
                0,
            )
        )

        status = self._to_text(
            value(
                "status",
                "В работе",
            )
        )

        if not status:
            status = "В работе"

        return {
            "model": self._to_text(
                value("model")
            ),
            "serial_number": self._to_text(
                value("serial_number")
            ),
            "garage_number": self._to_text(
                value("garage_number")
            ),
            "registration_number": self._to_text(
                value("registration_number")
            ),
            "manufacture_year": manufacture_year,
            "current_hours": current_hours,
            "status": status,
            "note": self._to_text(
                value("note")
            ),
        }

    def _validate_equipment_row(
        self,
        data,
        row_number,
    ):
        missing = []

        if not data["model"]:
            missing.append("Модель")

        if not data["serial_number"]:
            missing.append(
                "Серийный номер"
            )

        if not data["garage_number"]:
            missing.append(
                "Гаражный номер"
            )

        if missing:
            raise ValueError(
                "не заполнены обязательные поля: "
                + ", ".join(missing)
            )

    def _is_empty_row(self, data):
        return not any(
            (
                data["model"],
                data["serial_number"],
                data["garage_number"],
                data["registration_number"],
                data["manufacture_year"],
                data["current_hours"],
                data["note"],
            )
        )

    @staticmethod
    def _normalize_header(value):
        if value is None:
            return ""

        return " ".join(
            str(value)
            .strip()
            .lower()
            .replace("ё", "е")
            .split()
        )

    @staticmethod
    def _to_text(value):
        if value is None:
            return ""

        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))

        return str(value).strip()

    @staticmethod
    def _to_float(value):
        if value in (
            None,
            "",
        ):
            return 0.0

        if isinstance(
            value,
            (int, float),
        ):
            return float(value)

        text = (
            str(value)
            .strip()
            .replace(" ", "")
            .replace(",", ".")
        )

        try:
            return float(text)
        except ValueError:
            raise ValueError(
                f"некорректная наработка: {value}"
            )

    @staticmethod
    def _to_int_or_none(value):
        if value in (
            None,
            "",
        ):
            return None

        if isinstance(value, int):
            return value

        if isinstance(value, float):
            return int(value)

        text = str(value).strip()

        try:
            return int(float(text))
        except ValueError:
            raise ValueError(
                f"некорректный год выпуска: {value}"
            )
